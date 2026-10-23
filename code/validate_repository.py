#!/usr/bin/env python3

"""Validate final PDFs and the structure of Data S1."""

from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parent.parent


def validate_pdfs() -> None:
    main = sorted((ROOT / "figures").glob("Figure_*.pdf"))
    supplementary = sorted(
        (ROOT / "figures" / "supplementary").glob("Figure_S*.pdf")
    )
    expected = [f"Figure_{number}.pdf" for number in range(1, 6)]
    expected_supplementary = [f"Figure_S{number}.pdf" for number in range(1, 4)]
    assert [path.name for path in main] == expected, main
    assert [path.name for path in supplementary] == expected_supplementary, supplementary

    pdfs = sorted(ROOT.rglob("*.pdf"))
    for path in pdfs:
        reader = PdfReader(str(path), strict=True)
        assert len(reader.pages) >= 1, f"No pages found in {path}"
    print(f"Validated {len(pdfs)} PDFs, including 8 numbered figures.")


def validate_data_s1() -> None:
    path = ROOT / "data" / "supplementary" / "Data_S1_functional_gene_sets.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=False)
    assert workbook.sheetnames == ["README", "Depolymerisation", "Uptake"]

    depolymerisation = [
        row[1]
        for row in workbook["Depolymerisation"].iter_rows(
            min_row=2, values_only=True
        )
        if row[1] is not None
    ]
    uptake = [
        row[2]
        for row in workbook["Uptake"].iter_rows(min_row=2, values_only=True)
        if row[2] is not None
    ]
    assert len(depolymerisation) == len(set(depolymerisation)) == 41
    assert len(uptake) == len(set(uptake)) == 211
    print("Validated Data S1: 41 GH families and 211 uptake KOs.")


if __name__ == "__main__":
    validate_pdfs()
    validate_data_s1()
